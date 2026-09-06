from fastapi import FastAPI, UploadFile
from fastapi.responses import StreamingResponse, HTMLResponse
import openpyxl
import io

app = FastAPI()

@app.get("/")
async def form():
    html = """
    <html>
    <body>
      <h2>Gamesheet 自動生成</h2>
      <form action="/generate" method="post" enctype="multipart/form-data">
        <p>AllmemberA: <input type="file" name="allA"></p>
        <p>AllmemberB: <input type="file" name="allB"></p>
        <p>Checksheet: <input type="file" name="check"></p>
        <button type="submit">生成する</button>
      </form>
    </body>
    </html>
    """
    return HTMLResponse(content=html)

@app.post("/generate")
async def generate(allA: UploadFile, allB: UploadFile, check: UploadFile):
    wbA = openpyxl.load_workbook(allA.file)
    wbB = openpyxl.load_workbook(allB.file)
    wbC = openpyxl.load_workbook(check.file)
    wbG = openpyxl.load_workbook("Gamesheet.xlsx")

    wsG = wbG["Gamesheet"]
    wsA = wbA.active
    wsB = wbB.active
    wsC = wbC.active

    wsG["G4"] = wsA["B2"].value
    wsG["G29"] = wsB["B2"].value

    rowA, rowG = 20, 10
    while wsA[f"A{rowA}"].value:
        wsG[f"A{rowG}"] = wsA[f"A{rowA}"].value
        wsG[f"B{rowG}"] = wsA[f"B{rowA}"].value
        wsG[f"C{rowG}"] = wsA[f"C{rowA}"].value
        wsG[f"D{rowG}"] = wsA[f"D{rowA}"].value
        wsG[f"E{rowG}"] = wsA[f"E{rowA}"].value
        rowA += 1
        rowG += 1

    rowB, rowG = 20, 35
    while wsB[f"A{rowB}"].value:
        wsG[f"A{rowG}"] = wsB[f"A{rowB}"].value
        wsG[f"B{rowG}"] = wsB[f"B{rowB}"].value
        wsG[f"C{rowG}"] = wsB[f"C{rowB}"].value
        wsG[f"D{rowG}"] = wsB[f"D{rowB}"].value
        wsG[f"E{rowG}"] = wsB[f"E{rowB}"].value
        rowB += 1
        rowG += 1

    rowC, rowG = 5, 60
    while wsC[f"A{rowC}"].value:
        wsG[f"A{rowG}"] = wsC[f"A{rowC}"].value
        wsG[f"B{rowG}"] = wsC[f"B{rowC}"].value
        wsG[f"C{rowG}"] = wsC[f"C{rowC}"].value
        wsG[f"D{rowG}"] = wsC[f"D{rowC}"].value
        wsG[f"E{rowG}"] = wsC[f"E{rowC}"].value
        rowC += 1
        rowG += 1

    output = io.BytesIO()
    wbG.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=Gamesheet_completed.xlsx"}
    )
